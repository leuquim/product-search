"""
Enhanced database module with multi-file support
"""
import duckdb
import os
import time
from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime
from .config import DATABASE_PATH, DEFAULT_SEARCH_LIMIT


class MultiFileDatabase:
    def __init__(self, db_path: str = DATABASE_PATH):
        self.db_path = os.path.abspath(db_path)
        self.conn = None
        self._connect()
        self._ensure_schema()
    
    def _clean_column_name(self, col: str) -> str:
        """Clean column names consistently for table creation and inserts"""
        clean_col = str(col).strip()
        clean_col = clean_col.replace(' ', '_').replace('-', '_').replace('.', '_')
        clean_col = clean_col.replace('/', '_').replace('(', '').replace(')', '')
        clean_col = clean_col.replace('[', '').replace(']', '').replace(',', '_')
        clean_col = ''.join(c for c in clean_col if c.isalnum() or c == '_')
        if clean_col and clean_col[0].isdigit():
            clean_col = f"col_{clean_col}"
        if not clean_col:
            clean_col = "unknown_column"
        return clean_col

    def _connect(self):
        """Establish connection to DuckDB database"""
        try:
            self.conn = duckdb.connect(self.db_path)
            print(f"Connected to database: {self.db_path}")
        except Exception as e:
            raise Exception(f"Failed to connect to database: {e}")

    def _ensure_schema(self):
        """Ensure all required tables exist"""
        try:
            # Run migration to ensure schema is up to date
            from .migrations.migration_001_multi_file_support import upgrade
            upgrade(self.db_path)
        except Exception as e:
            print(f"Migration check failed: {e}")
            # Create basic schema if migration fails
            self._create_basic_schema()

    def _create_basic_schema(self):
        """Create basic schema if migration fails"""
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS imported_files (
                file_id INTEGER PRIMARY KEY,
                original_filename VARCHAR NOT NULL,
                import_date TIMESTAMP NOT NULL,
                row_count INTEGER DEFAULT 0,
                indexed_columns VARCHAR,
                file_size_mb DOUBLE,
                status VARCHAR DEFAULT 'active'
            )
        """)
        
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS file_columns (
                file_id INTEGER,
                column_name VARCHAR,
                is_indexed BOOLEAN DEFAULT FALSE,
                data_type VARCHAR,
                PRIMARY KEY (file_id, column_name)
            )
        """)

    def close(self):
        """Close database connection"""
        if self.conn:
            self.conn.close()
            self.conn = None

    def get_next_file_id(self) -> int:
        """Get the next available file ID"""
        result = self.conn.execute("""
            SELECT COALESCE(MAX(file_id), 0) + 1 FROM imported_files
        """).fetchone()
        return result[0]

    def register_file(self, filename: str, row_count: int = 0, 
                     indexed_columns: List[str] = None, file_size_mb: float = 0) -> int:
        """Register a new imported file and return its ID"""
        file_id = self.get_next_file_id()
        
        indexed_cols_str = ','.join(indexed_columns) if indexed_columns else ''
        
        self.conn.execute("""
            INSERT INTO imported_files 
            (file_id, original_filename, import_date, row_count, indexed_columns, file_size_mb, status)
            VALUES (?, ?, CURRENT_TIMESTAMP, ?, ?, ?, 'active')
        """, [file_id, filename, row_count, indexed_cols_str, file_size_mb])
        
        return file_id

    def update_file_stats(self, file_id: int, row_count: int):
        """Update file statistics after import"""
        self.conn.execute("""
            UPDATE imported_files 
            SET row_count = ?, updated_at = CURRENT_TIMESTAMP
            WHERE file_id = ?
        """, [row_count, file_id])

    def register_columns(self, file_id: int, columns: List[str], indexed_columns: List[str]):
        """Register columns for a file"""
        for col in columns:
            clean_col = self._clean_column_name(col)
            is_indexed = col in indexed_columns
            
            self.conn.execute("""
                INSERT INTO file_columns (file_id, column_name, is_indexed)
                VALUES (?, ?, ?)
                ON CONFLICT (file_id, column_name) 
                DO UPDATE SET is_indexed = EXCLUDED.is_indexed
            """, [file_id, clean_col, is_indexed])

    def create_products_table_if_not_exists(self, columns: List[str]):
        """Create products table if it doesn't exist"""
        table_exists = self.conn.execute("""
            SELECT COUNT(*) FROM information_schema.tables 
            WHERE table_name = 'products'
        """).fetchone()[0] > 0
        
        if not table_exists:
            column_defs = []
            column_defs.append("source_file_id INTEGER")
            
            for col in columns:
                clean_col = self._clean_column_name(col)
                column_defs.append(f'"{clean_col}" VARCHAR')
            
            create_sql = f"""
            CREATE TABLE products (
                {', '.join(column_defs)}
            )
            """
            
            self.conn.execute(create_sql)
            print(f"Created products table with {len(columns)} columns")

    def ensure_columns_exist(self, columns: List[str]):
        """Ensure all columns exist in products table"""
        existing_columns = self.conn.execute("""
            SELECT column_name 
            FROM information_schema.columns 
            WHERE table_name = 'products'
        """).fetchall()
        
        existing_set = {col[0].lower() for col in existing_columns}
        
        for col in columns:
            clean_col = self._clean_column_name(col)
            if clean_col.lower() not in existing_set:
                try:
                    self.conn.execute(f"""
                        ALTER TABLE products ADD COLUMN "{clean_col}" VARCHAR
                    """)
                    print(f"Added new column: {clean_col}")
                except Exception as e:
                    print(f"Could not add column {clean_col}: {e}")

    def insert_batch(self, data: List[Dict[str, Any]], file_id: int) -> int:
        """Insert batch of records with file ID - optimized version"""
        if not data:
            return 0

        columns = list(data[0].keys())
        self.create_products_table_if_not_exists(columns)
        self.ensure_columns_exist(columns)
        
        clean_columns = [self._clean_column_name(col) for col in columns]
        
        # Add source_file_id to the insert
        all_columns = ['source_file_id'] + clean_columns
        placeholders = ', '.join(['?' for _ in all_columns])
        
        insert_sql = f"""
        INSERT INTO products ({', '.join([f'"{col}"' for col in all_columns])})
        VALUES ({placeholders})
        """

        try:
            # Prepare batch data more efficiently
            batch_data = []
            for row in data:
                # Avoid string conversion unless necessary
                row_values = [file_id]
                for col in columns:
                    val = row.get(col)
                    if val is None:
                        row_values.append('')
                    elif isinstance(val, str):
                        row_values.append(val)
                    else:
                        row_values.append(str(val))
                batch_data.append(row_values)
            
            # Use transaction for better performance
            self.conn.begin()
            self.conn.executemany(insert_sql, batch_data)
            self.conn.commit()
            
            return len(batch_data)
        except Exception as e:
            self.conn.rollback()
            raise Exception(f"Failed to insert batch: {e}")

    def create_indexes(self, file_id: int, columns: List[str]):
        """Create indexes on specified columns"""
        for col in columns:
            clean_col = self._clean_column_name(col)
            try:
                index_name = f"idx_{clean_col.lower()}_{file_id}"
                self.conn.execute(f"""
                    CREATE INDEX IF NOT EXISTS {index_name} 
                    ON products ("{clean_col}") 
                    WHERE source_file_id = {file_id}
                """)
                print(f"Created index on {clean_col} for file {file_id}")
            except Exception as e:
                print(f"Could not create index on {clean_col}: {e}")

    def search(self, query: str, file_ids: List[int] = None, 
              limit: int = DEFAULT_SEARCH_LIMIT, offset: int = 0) -> Tuple[List[Dict], int]:
        """Search across specified files or all files"""
        if not self._table_exists('products'):
            return [], 0

        start_time = time.time()
        
        # Get indexed columns for the files we're searching
        if file_ids:
            file_filter = f"WHERE file_id IN ({','.join(map(str, file_ids))}) AND is_indexed = true"
        else:
            file_filter = "WHERE is_indexed = true"
        
        indexed_cols = self.conn.execute(f"""
            SELECT DISTINCT column_name 
            FROM file_columns 
            {file_filter}
        """).fetchall()
        
        if not indexed_cols:
            # Fallback to default columns if none are indexed
            indexed_cols = [('ASSEMBLY',), ('DESCRIPTION',)]
        
        # Build search query
        conditions = []
        params = []
        
        if query.strip():
            for col_tuple in indexed_cols:
                col = col_tuple[0]
                conditions.append(f'"{col}" ILIKE ?')
                params.append(f'%{query}%')
        
        where_clause = ""
        if conditions:
            where_clause = f"WHERE ({' OR '.join(conditions)})"
        
        if file_ids:
            file_condition = f"source_file_id IN ({','.join(map(str, file_ids))})"
            if where_clause:
                where_clause += f" AND {file_condition}"
            else:
                where_clause = f"WHERE {file_condition}"
        
        # Count total results
        count_sql = f"""
            SELECT COUNT(*) FROM products {where_clause}
        """
        total_count = self.conn.execute(count_sql, params).fetchone()[0]
        
        # Get paginated results with file info
        search_sql = f"""
            SELECT p.*, f.original_filename 
            FROM products p
            LEFT JOIN imported_files f ON p.source_file_id = f.file_id
            {where_clause}
            ORDER BY p.source_file_id, p.rowid
            LIMIT ? OFFSET ?
        """
        
        params.extend([limit, offset])
        results = self.conn.execute(search_sql, params).fetchall()
        
        columns = [desc[0] for desc in self.conn.description]
        formatted_results = []
        
        for row in results:
            row_dict = dict(zip(columns, row))
            # Move source filename to a special key
            if 'original_filename' in row_dict:
                row_dict['_source_file'] = row_dict.pop('original_filename')
            formatted_results.append(row_dict)
        
        search_time = time.time() - start_time
        print(f"Search completed in {search_time:.3f}s - Found {total_count} matches")
        
        return formatted_results, total_count

    def search_grouped(self, query: str, file_ids: List[int] = None,
                      limit_per_file: int = 10) -> Dict[str, Any]:
        """Search and group results by source file"""
        if not self._table_exists('products'):
            return {"files": {}, "total_count": 0}
        
        # Get list of files to search
        if file_ids:
            files = self.conn.execute("""
                SELECT file_id, original_filename, row_count 
                FROM imported_files 
                WHERE file_id IN ({}) AND status = 'active'
            """.format(','.join(map(str, file_ids)))).fetchall()
        else:
            files = self.conn.execute("""
                SELECT file_id, original_filename, row_count 
                FROM imported_files 
                WHERE status = 'active'
            """).fetchall()
        
        grouped_results = {}
        total_count = 0
        
        for file_id, filename, row_count in files:
            results, count = self.search(query, [file_id], limit_per_file, 0)
            if count > 0:
                grouped_results[filename] = {
                    "file_id": file_id,
                    "results": results,
                    "count": count,
                    "total_rows": row_count
                }
                total_count += count
        
        return {
            "files": grouped_results,
            "total_count": total_count,
            "file_count": len(grouped_results)
        }

    def get_imported_files(self) -> List[Dict]:
        """Get list of all imported files"""
        results = self.conn.execute("""
            SELECT file_id, original_filename, import_date, row_count, 
                   indexed_columns, file_size_mb, status
            FROM imported_files
            ORDER BY import_date DESC
        """).fetchall()
        
        files = []
        for row in results:
            files.append({
                "file_id": row[0],
                "filename": row[1],
                "import_date": row[2],
                "row_count": row[3],
                "indexed_columns": row[4].split(',') if row[4] else [],
                "file_size_mb": row[5],
                "status": row[6]
            })
        
        return files

    def delete_file(self, file_id: int) -> bool:
        """Delete a file and its associated data"""
        try:
            # Delete products from this file
            self.conn.execute("""
                DELETE FROM products WHERE source_file_id = ?
            """, [file_id])
            
            # Delete column info
            self.conn.execute("""
                DELETE FROM file_columns WHERE file_id = ?
            """, [file_id])
            
            # Mark file as deleted (soft delete)
            self.conn.execute("""
                UPDATE imported_files 
                SET status = 'deleted', updated_at = CURRENT_TIMESTAMP
                WHERE file_id = ?
            """, [file_id])
            
            self.conn.commit()
            return True
        except Exception as e:
            print(f"Error deleting file {file_id}: {e}")
            self.conn.rollback()
            return False

    def get_stats(self) -> Dict[str, Any]:
        """Get database statistics"""
        stats = {
            "total_files": 0,
            "total_records": 0,
            "database_size_mb": 0,
            "files": []
        }
        
        try:
            # Get file count and stats
            file_stats = self.conn.execute("""
                SELECT COUNT(*), SUM(row_count), SUM(file_size_mb)
                FROM imported_files
                WHERE status = 'active'
            """).fetchone()
            
            stats["total_files"] = file_stats[0] or 0
            stats["total_records"] = file_stats[1] or 0
            stats["total_file_size_mb"] = file_stats[2] or 0
            
            # Get database file size
            if os.path.exists(self.db_path):
                stats["database_size_mb"] = round(os.path.getsize(self.db_path) / (1024 * 1024), 2)
            
            # Get per-file stats
            stats["files"] = self.get_imported_files()
            
        except Exception as e:
            print(f"Error getting stats: {e}")
        
        return stats

    def _table_exists(self, table_name: str) -> bool:
        """Check if a table exists"""
        try:
            result = self.conn.execute("""
                SELECT COUNT(*) FROM information_schema.tables 
                WHERE table_name = ?
            """, [table_name]).fetchone()
            return result[0] > 0
        except:
            return False

    def preview_excel_data(self, file_path: str, rows: int = 10) -> Tuple[List[Dict], List[str]]:
        """Preview Excel file data without importing"""
        import openpyxl
        
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet = workbook.active
            
            # Get headers
            headers = []
            for cell in sheet[1]:
                if cell.value is not None and str(cell.value).strip():
                    headers.append(str(cell.value).strip())
                else:
                    headers.append(f"Column_{len(headers)+1}")
            
            # Get sample data
            preview_data = []
            row_count = 0
            
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if row_count >= rows:
                    break
                
                row_dict = {}
                has_data = False
                
                for idx, value in enumerate(row):
                    if idx < len(headers):
                        if value is not None:
                            row_dict[headers[idx]] = str(value).strip()
                            has_data = True
                        else:
                            row_dict[headers[idx]] = ""
                
                if has_data:
                    preview_data.append(row_dict)
                    row_count += 1
            
            workbook.close()
            return preview_data, headers
            
        except Exception as e:
            raise Exception(f"Failed to preview file: {e}")

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()